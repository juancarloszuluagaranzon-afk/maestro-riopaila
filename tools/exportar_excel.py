"""Convierte maestro.csv a un .xlsx para repartir fuera de la app.

Nace de la suspensión del 19-oct-2026: con GitHub Pages apagado, el campo no
puede consultar el Maestro y necesita los datos por otra vía.

Dos hojas, el mismo alcance que usa la app: Riopaila (EMPRESA = RIOP) y
Castilla (CAST + CAUC). Una tercera, "Acerca de", dice de cuándo son los datos.

OJO CON LA EDAD. La columna `EDAD HOY MESES` del CSV NO se mantiene al día:
`maestro.html` la recalcula cada vez que carga el archivo, así que en el repo
queda congelada en la fecha de la última importación. Un Excel no recalcula
nada, así que aquí se vuelve a calcular **con la fecha de generación** y con la
misma fórmula de la app: RENOVACION da 0; si no, desde la más reciente entre la
siembra y el último corte, dividido por 365,25/12. Sin esto, el archivo repartido
diría que una suerte cortada ayer tiene doce meses.

Uso (desde la raíz del repo):
  python tools/exportar_excel.py                      # -> maestro_YYYY-MM-DD.xlsx
  python tools/exportar_excel.py --salida "C:/ruta/archivo.xlsx"
"""
import argparse
import datetime as dt
import io
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

DIAS_POR_MES = 365.25 / 12
NUMERICAS = {
    "AREA NETA HA", "NUMERO DE CORTE", "TONELADAS PPTO", "TCH PPTO",
    "TCH ANTERIOR", "TCHM ANTERIOR", "TCH RANDOM FOREST", "TONELADAS ESTIMADAS",
}
FECHAS = {"FECHA DE SIEMBRA", "FECHA DE ULTIMO CORTE", "FECHA DEL PROXIMO CORTE"}
DECIMALES = {"AREA NETA HA": "0.00", "EDAD HOY MESES": "0.00", "TCH ANTERIOR": "0.0",
             "TCHM ANTERIOR": "0.00", "TCH RANDOM FOREST": "0.0"}
# El ingenio exporta a veces la fecha como serial de Excel mezclado con texto.
EPOCA_EXCEL = dt.date(1899, 12, 30)


def leer_csv(ruta):
    texto = io.open(ruta, encoding="utf-8-sig", newline="").read()
    filas = [l for l in texto.replace("\r\n", "\n").split("\n") if l.strip()]
    cab = filas[0].split(";")
    return cab, [l.split(";") for l in filas[1:]]


def fecha(valor):
    """'5/10/2026', un serial de Excel o vacío -> date | None."""
    s = (valor or "").strip()
    if not s:
        return None
    m = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        d, mes, a = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return dt.date(a, mes, d)
        except ValueError:
            return None
    if re.fullmatch(r"\d{4,5}", s) and 20000 <= int(s) <= 60000:
        return EPOCA_EXCEL + dt.timedelta(days=int(s))
    return None


def numero(valor):
    s = (valor or "").strip()
    if not s:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def edad_meses(variedad, siembra, ultimo, hoy):
    """La fórmula de la app (lib/maestro/sincronizar.ts y maestro.html)."""
    if re.search(r"RENOVACI[OÓ]N", variedad or "", re.I):
        return 0.0
    fs, fc = fecha(siembra), fecha(ultimo)
    base = max(x for x in (fs, fc) if x is not None) if (fs or fc) else None
    if base is None:
        return None
    return round(max(0, (hoy - base).days) / DIAS_POR_MES, 2)


def hoja(wb, titulo, cab, filas, hoy):
    ws = wb.create_sheet(titulo)
    ws.append(cab)
    i_var, i_sie, i_ult = (cab.index(c) for c in ("VARIEDAD", "FECHA DE SIEMBRA", "FECHA DE ULTIMO CORTE"))
    i_edad = cab.index("EDAD HOY MESES")

    for f in filas:
        salida = []
        for j, col in enumerate(cab):
            bruto = f[j] if j < len(f) else ""
            if col == "EDAD HOY MESES":
                salida.append(edad_meses(f[i_var], f[i_sie], f[i_ult], hoy))
            elif col in FECHAS:
                salida.append(fecha(bruto) or (bruto.strip() or None))
            elif col in NUMERICAS:
                salida.append(numero(bruto))
            elif col == "ZONA":
                n = numero(bruto)
                salida.append(int(n) if n is not None and n == int(n) else (bruto.strip() or None))
            else:
                salida.append(bruto.strip() or None)
        ws.append(salida)

    cabecera = Font(bold=True, color="FFFFFF")
    fondo = PatternFill("solid", fgColor="0F766E")
    for celda in ws[1]:
        celda.font, celda.fill = cabecera, fondo
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    for j, col in enumerate(cab, start=1):
        letra = get_column_letter(j)
        ws.column_dimensions[letra].width = max(10, min(28, len(col) + 2))
        fmt = DECIMALES.get(col) or ("dd/mm/yyyy" if col in FECHAS else None)
        if fmt:
            for celda in ws[letra][1:]:
                celda.number_format = fmt
    return ws


def main():
    p = argparse.ArgumentParser()
    p.add_argument("--csv", default="maestro.csv")
    p.add_argument("--salida")
    args = p.parse_args()

    hoy = dt.date.today()
    cab, filas = leer_csv(args.csv)
    i_emp = cab.index("EMPRESA")
    riop = [f for f in filas if f[i_emp].strip() == "RIOP"]
    cast = [f for f in filas if f[i_emp].strip() in ("CAST", "CAUC")]
    otras = len(filas) - len(riop) - len(cast)

    wb = Workbook()
    wb.remove(wb.active)
    hoja(wb, "Riopaila", cab, riop, hoy)
    hoja(wb, "Castilla", cab, cast, hoy)

    ws = wb.create_sheet("Acerca de")
    for linea in [
        ["Maestro de Suertes — Riopaila / Castilla"],
        [],
        ["Generado", hoy.strftime("%d/%m/%Y")],
        ["Suertes", f"{len(riop)} Riopaila · {len(cast)} Castilla (CAST + CAUC) · {len(filas)} en total"],
        [],
        ["EDAD HOY MESES está calculada al día de generación, desde la más reciente"],
        ["entre FECHA DE SIEMBRA y FECHA DE ULTIMO CORTE. Las suertes en RENOVACION"],
        ["quedan en 0. En un Excel ese dato NO se actualiza solo: vuelve a exportar"],
        ["el archivo cuando necesites la edad al día."],
        [],
        ["TCH ANTERIOR viene de un reporte aparte del ingenio y puede ir por detrás"],
        ["de los cortes más recientes. TCH RANDOM FOREST queda vacío en las suertes"],
        ["que acaban de cortarse: ese estimado era de la caña ya cosechada."],
    ]:
        ws.append(linea)
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 72
    ws.column_dimensions["B"].width = 60

    salida = args.salida or f"maestro_{hoy.isoformat()}.xlsx"
    wb.save(salida)
    print(f"{salida}")
    print(f"  Riopaila {len(riop)} · Castilla {len(cast)}" + (f" · otras empresas {otras}" if otras else ""))
    print(f"  EDAD HOY MESES recalculada al {hoy.strftime('%d/%m/%Y')}")


if __name__ == "__main__":
    main()
